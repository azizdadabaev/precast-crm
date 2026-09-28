# Read the live Donabay.xlsx into the build pack's seed format (Spec 06) so the
# CRM import always uses the latest data. Adds the owner-approved DEBT_CAP 50 %
# from 2026-09-28 and sets every joining date to 2026-08-31 (spec §3).
# Usage: python scripts/crew-pay/extract-donabay.py path/to/Donabay.xlsx out.json
import json, sys, datetime as dt
import openpyxl

FIRST = "2026-08-31"
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
iso = lambda v: v.date().isoformat() if isinstance(v, dt.datetime) else (v.isoformat() if isinstance(v, dt.date) else None)

ws = wb["Workers"]
workers, by_name, slot_code = [], {}, {}
for r in range(6, 18):
    code, name, status = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
    slot_code[r - 6] = code
    if not name:
        continue
    left = iso(ws.cell(r, 5).value)
    workers.append({"code": code, "name": str(name).strip(), "status": "left" if status == "Left" else "active",
                    "joined_on": FIRST, "left_on": left, "phone": ws.cell(r, 6).value, "notes": ws.cell(r, 7).value})
    by_name[str(name).strip()] = code

ws = wb["Daily Log"]
days = []
for r in range(7, 378):
    date = iso(ws.cell(r, 1).value)
    if not date:
        continue
    moulded, broken = ws.cell(r, 3).value, ws.cell(r, 4).value
    att = {}
    for k in range(12):                       # G..R = worker slots W01..W12
        v = ws.cell(r, 7 + k).value
        if isinstance(v, (int, float)) and v in (0.5, 1):
            att[slot_code[k]] = v
    if moulded is None and not att and broken in (None, 0):
        continue
    days.append({"date": date, "moulded": moulded, "broken": broken or 0, "attendance": att, "notes": ws.cell(r, 24).value})

TYPES = {"Advance": "advance", "Weekly pay": "weekly_pay", "Correction": "correction"}
ws = wb["Cash Ledger"]
ledger, seq = [], 0
for r in range(7, 1007):
    date, name, typ, amount = iso(ws.cell(r, 2).value), ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value
    if not date:
        continue
    seq += 1
    reason = ws.cell(r, 6).value
    method = "offset" if reason and "NOT cash" in str(reason) else "cash"
    ledger.append({"seq": seq, "date": date, "worker": by_name[str(name).strip()], "type": TYPES[typ],
                   "amount": int(round(amount)), "reason": reason, "given_by": ws.cell(r, 7).value,
                   "signed": ws.cell(r, 8).value == "Yes", "method": method})

ws = wb["Settings"]
settings = [
    {"key": "break_allowance_pct", "value": float(ws["C6"].value or 0), "effective_from": FIRST},
    {"key": "debt_cap_pct", "value": float(ws["C7"].value), "effective_from": FIRST},
    {"key": "reject_alarm_pct", "value": float(ws["C8"].value), "effective_from": FIRST},
    {"key": "debt_cap_pct", "value": 0.5, "effective_from": "2026-09-28"},
]
rates = [{"effective_from": iso(ws.cell(r, 2).value), "rate_per_block": int(ws.cell(r, 3).value)}
         for r in range(13, 33) if iso(ws.cell(r, 2).value) and ws.cell(r, 3).value]

seed = {"_source": sys.argv[1], "currency": "UZS", "timezone": "Asia/Tashkent", "first_week_start": FIRST,
        "settings": settings, "pay_rates": rates, "workers": workers, "daily_log": days, "cash_ledger": ledger}
with open(sys.argv[2], "w", encoding="utf-8") as f:
    json.dump(seed, f, indent=1, ensure_ascii=False, default=str)
print(f"workers {len(workers)} · days {len(days)} · ledger {len(ledger)} · rates {len(rates)}")
